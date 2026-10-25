"""
요구사항 정의서(엑셀 양식) 생성 — 10/26 분석 보고 제출용, Problem Description 범위
  .venv/bin/python scripts/build_requirements_xlsx.py <팀원 초안.xlsx> docs/01-1_요구사항정의서.xlsx

* 팀원 초안(양식)을 복사해 서식·드롭다운·병합 셀을 그대로 유지하고 내용만 보강한다.
* '결정 내용'(노란 칸)에는 팀이 확정한 가정을 적는다 (구현과 일치. 구현하지 않기로 한 것도 이유와 함께 적음).
"""
import sys
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)

# ---------------------------------------------------------------- 시트 1: 요구사항 정의
ws = wb["요구사항정의_양식"]
STYLE_ROW = 4          # 데이터 행 서식 견본

def style_row(r):
    for col in "ABCDE":
        s, d = ws[f"{col}{STYLE_ROW}"], ws[f"{col}{r}"]
        d.font, d.border, d.alignment, d.fill, d.number_format = copy(s.font), copy(s.border), copy(s.alignment), copy(s.fill), s.number_format

def lines(text):
    return sum(max(1, -(-len(line) // 48)) for line in text.split("\n"))   # C열 85폭 ≈ 한글 48자

def put(r, rid, name, items, kind, entities):
    style_row(r)
    body = "\n".join(f"{i}. {t}" for i, t in enumerate(items, 1))
    ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = rid, name, body, kind, entities
    ws.row_dimensions[r].height = max(30, lines(body) * 12.5 + 8)

# 팀원 초안 보정 (내용은 유지하고 용어·유형만 정리)
assert ws["E4"].value == "직원, 개발자, 일반직원"
ws["E4"] = "직원, 개발자, 지원직"
ws["C4"] = ws["C4"].value.replace("그 외 직원(마케팅, 경영관리, 연구개발)", "지원직(마케팅, 경영관리, 연구개발 직원)")
# '별도로'는 핵심 가이드라인(발주처 별도 엔티티)의 표현이라 10/26 범위(Problem Description only)에서 뺌
OLD_URQ004 = "1. 프로젝트를 발주한 고객사(발주처) 정보를 프로젝트와 별도로 관리해야 한다."
assert OLD_URQ004 in ws["C7"].value
ws["C7"] = ws["C7"].value.replace(OLD_URQ004, "1. 프로젝트를 발주한 고객사(발주처) 정보를 관리해야 한다.")

# URQ009: 화면 구성은 '디자인', 웹 기반 시스템 항목은 URQ011(시스템)로 이동
put(12, "URQ009", "웹 페이지\n(화면 구성)", [
    "회원가입 페이지에는 직원 정보 입력 폼과 ID 중복 확인 버튼이 있다.",
    "ID 중복 확인 결과(사용 가능 / 이미 등록됨)를 화면에 바로 표시한다.",
    "로그인한 직원은 자신의 정보를 수정하는 페이지로 이동할 수 있다.",
    "경영진으로 로그인하면 직원 검색, 프로젝트 현황, 투입 현황 조회 메뉴를 사용할 수 있다.",
    "프로젝트 종료 후 참여 직원은 평가 입력 화면에서 항목별 평점과 평가 내용을 입력한다.",
], "디자인", "직원, 프로젝트, 프로젝트투입, 평가")

# 10/26 은 "Problem Description only" — 개인정보 노출·비밀번호 저장 방식, 시스템 환경·데이터 설계(과제 2절 Requirements)는 넣지 않는다
put(13, "URQ010", "권한 관리", [
    "일반 직원은 본인 정보만 조회·수정하고, 경영진만 타 직원 정보를 검색·조회한다.",
    "PM 권한(PM 평가, 고객 평가 의뢰·보고)은 해당 프로젝트에 PM 으로 투입된 직원에게만 준다.",
], "정책", "직원, 경영진, 프로젝트투입, 평가")

# 출처 메모 (양식 아래 빈 행)
ws["A15"] = ("※ URQ001~URQ010 은 모두 Problem Description 문장에서 도출 (과제 일정표의 '요구사항 정의서 (“Problem Description” only!)' 에 따름). "
             "문장별 대응은 '추적표' 시트 참고.")
ws.merge_cells("A15:E15")
ws["A15"].font = Font(name="돋움", size=9, color="666666")
ws["A15"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[15].height = 28

# ---------------------------------------------------------------- 시트 2: 팀 확인 필요 사항
q = wb["팀 확인 필요 사항"]
HDR, ROW, DEC = q["D4"], q["B5"], q["D5"]          # 헤더 / 일반 셀 / 결정(노란) 셀 견본

def clone(dst, src_cell):
    dst.font, dst.border, dst.alignment, dst.fill = copy(src_cell.font), copy(src_cell.border), copy(src_cell.alignment), copy(src_cell.fill)

assert q["A2"].value.startswith("※ 회의에서")
q["A2"] = "※ 문제 설명에 없는 부분은 아래 '결정 내용'과 같이 가정하여 정함. 결정 내용은 11/09 설계 보고의 가정·비즈니스 로직에 반영"
q.column_dimensions["D"].width = 60

decisions = {
    5:  "E + 입사연도 2자리 + 연도별 순번 4자리 (예: E260005). 동시 가입에도 중복되지 않게 DB 가 채번",
    6:  "회원가입은 전 직원 가능. 프로젝트 투입은 개발자만 가능",
    7:  "직급이 아니라 계정에 별도 권한(경영진)을 부여해 구분",
    8:  "DB 에 저장하되 본인 외 화면에서는 뒷자리 마스킹. 실제 운영 시에는 개인정보보호법에 따라 컬럼 암호화"
        "(검색용 해시 별도)를 적용하며, 본 과제 데모에서는 생략",
    9:  "프로젝트당 발주처 1곳",
    10: "가능. 같은 기간 투입률 합계 100% 이하로 제한",
    11: "기간을 나눠 다른 직무로 다시 투입하는 방식으로 표현 (투입 이력 보존)",
    12: "포함. 개발 경력 기산일(입사 전 경력 포함)로 총 경력을 계산",
    13: "초급 / 중급 / 고급 / 특급 4단계. 경력 기준 권고 등급은 자동 계산, 확정은 경영진",
    14: "1~5점 정수",
    15: "대상. 해당 프로젝트에 한 번이라도 투입된 직원은 평가 대상",
    16: "로그인 없이 PM 이 보낸 전용 링크로 고객 담당자가 직접 입력",
    17: "PM 은 PM 평가 대상이 아님 (시나리오 정의상 PM 평가는 PM 이 팀원을 평가하는 것). "
        "PM 은 고객·동료 평가를 받으며, 종합점수는 받은 평가 유형만으로 산출",
}
for r, text in decisions.items():
    q[f"D{r}"] = text

extra = [  # (확인할 내용, 관련 요구사항, 결정 내용)
    ("프로젝트번호도 자동으로 부여하는가, 직접 입력하는가?", "URQ003",
     "P + 착수연도 + 순번 3자리 (예: P2026-011) 자동 부여"),
    ("기술 등급에 따라 맡을 수 있는 직무를 제한하는가? (예: PM 은 고급 이상)", "URQ005, URQ006",
     "PM 고급 이상, PL·분석자·설계자 중급 이상, 프로그래머·테스터 초급 이상"),
    ("한 프로젝트의 PM 은 한 명인가? 프로젝트 중간에 바뀔 수 있는가?", "URQ005, URQ007",
     "같은 기간에는 1명. 기간을 나눠 교체 가능"),
    ("기술 등급은 누가 바꾸며, 변경 이력을 남기는가?", "URQ006",
     "경영진만 변경, 변경 전후 등급·처리자·사유를 이력으로 자동 기록"),
    ("발주처에서 관리할 정보는 무엇이고, 한 발주처에 담당자가 여러 명일 수 있는가?", "URQ004",
     "발주처명·사업자등록번호·업종·주소·연락처 관리, 담당자는 여러 명"),
    ("평가 항목(업무수행, 커뮤니케이션)이 나중에 늘어날 수 있는가?", "URQ007",
     "평가 항목을 별도 목록으로 관리해 항목 추가 가능"),
    ("인센티브 기초자료는 어떤 형태이며, 고객·PM·동료 평가를 어떻게 합산하는가?", "URQ008",
     "고객 40% · PM 40% · 동료 20% 가중 평균으로 종합점수를 산출하고 S/A/B/C 등급 부여"),
]
r0 = 18
for i, (text, rel, dec) in enumerate(extra):
    r = r0 + i
    q[f"A{r}"], q[f"B{r}"], q[f"C{r}"], q[f"D{r}"] = 14 + i, text, rel, dec
    clone(q[f"A{r}"], q["A5"]); clone(q[f"B{r}"], ROW); clone(q[f"C{r}"], q["C5"]); clone(q[f"D{r}"], DEC)
for r in range(5, r0 + len(extra)):
    q[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    q.row_dimensions[r].height = max(30, -(-len(q[f"D{r}"].value) // 34) * 13 + 6)   # D열 60폭 ≈ 한글 34자
assert all(q[f"D{r}"].value for r in range(5, r0 + len(extra)))

# ---------------------------------------------------------------- 시트 3: 추적표
t = wb.create_sheet("추적표")
t["A1"] = "요구사항 추적표 (Problem Description 문장 → 요구사항 ID)"
t["A1"].font = copy(q["A1"].font)
t["A2"] = "¶1 회사와 투입 · ¶2 관리 정보와 회원가입 · ¶3 경영진 조회 · ¶4 평가 개요 · ¶5 평가 종류별 설명"
t["A2"].font = copy(q["A2"].font)
heads = ["문단", "Problem Description 원문 요지", "요구사항 ID"]
for i, h in enumerate(heads):
    c = t.cell(row=4, column=i + 1, value=h); clone(c, HDR)
rows = [
    ("¶1", "100명 중 개발자 70명이 월평균 10개 정도의 프로젝트를 수행", "URQ001, URQ003"),
    ("¶1", "프로젝트 초기부터 종료 시까지 또는 일정 기간만 투입", "URQ005"),
    ("¶1", "경력과 기술 등급에 따라 PM, PL, 분석자, 설계자, 프로그래머, 테스터 직무", "URQ005, URQ006"),
    ("¶2", "프로젝트번호, 프로젝트명, 착수일자/종료일자, 발주처 관리", "URQ003, URQ004"),
    ("¶2", "직원번호, 직원명, 주민등록번호, 최종학력 관리", "URQ001"),
    ("¶2", "투입 직원의 경력과 경험한 기술(Skill Set) 관리", "URQ006"),
    ("¶2", "회원가입 페이지로 자신의 정보를 저장하고 수정/변경", "URQ001, URQ002, URQ009"),
    ("¶2", "직원 사번은 정해진 규칙에 따라 자동 부여", "URQ001"),
    ("¶2", "로그인 ID 중복 체크, 중복 시 에러 메시지 후 새 ID 입력", "URQ002, URQ009"),
    ("¶3", "현재 몇 개의 프로젝트를 수행 중인지, 프로젝트별 투입 인원·직무·투입 기간", "URQ008"),
    ("¶3", "과거 특정 시점에 어떤 직원이 어느 프로젝트에 어떤 직무로 참여했는지", "URQ005, URQ008"),
    ("¶3", "개인별 경력관리와 인센티브 지급을 위한 기초자료 추출", "URQ006, URQ008"),
    ("¶3", "경영진은 타 직원 정보를 검색할 수 있는 권한", "URQ008, URQ010"),
    ("¶4", "프로젝트 종료 시점 평가: 고객 평가, PM 평가, 동료 평가", "URQ007"),
    ("¶4", "평가자와 피평가자, 평가 항목(업무수행, 커뮤니케이션)별 평점과 평가 내용", "URQ007"),
    ("¶5", "고객 평가: PM 이 주관하여 고객사 담당자에게 의뢰하고 결과를 회사에 보고", "URQ004, URQ007, URQ010"),
    ("¶5", "동료 평가: 자기 이외의 참여자 평가 / PM 평가: PM 이 팀원 평가", "URQ007"),
    ("¶5", "평가 결과는 인사고과와 인사 평가의 근거 자료", "URQ007, URQ008"),
]
for i, (p, txt, ids) in enumerate(rows):
    r = 5 + i
    for j, v in enumerate((p, txt, ids)):
        c = t.cell(row=r, column=j + 1, value=v)
        clone(c, q["A5"] if j != 1 else ROW)
    t.row_dimensions[r].height = 24
t.column_dimensions["A"].width, t.column_dimensions["B"].width, t.column_dimensions["C"].width = 10, 80, 26

wb.save(dst)
print("saved", dst)
